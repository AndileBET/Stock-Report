from datetime import date, datetime
from decimal import Decimal
from io import BytesIO

from openpyxl import load_workbook
from docx import Document

from backend.app.api import _build_executive_summary, _selection_meta
from backend.app.domain.models import StockTransaction
from backend.app.reporting.executive_report import (
    build_docx_report,
    build_pdf_report,
    build_text_report,
)
from backend.app.reporting.report_service import save_transactions
from backend.app.reporting.stock_table_export import build_docx_table, build_pdf_table, build_xlsx_table


def make_tx(region, team, item_code, item_name, tx_type, qty, value, job, dt):
    return StockTransaction(
        date=dt,
        job_reference=job,
        source_site="",
        destination_site=f"{region} Team",
        operator="Tester",
        classification="Stock",
        item_code=item_code,
        item_name=item_name,
        category="HWM - Stock",
        sub_category="",
        unit="Unit",
        transaction_type=tx_type,
        quantity=Decimal(str(qty)),
        value=Decimal(str(value)),
        batch="B1",
        actioned=True,
        region=region,
        team=team,
    )


def test_executive_summary_supports_multiple_regions_and_teams(monkeypatch):
    txs = [
        make_tx("FS", "FS Team 07", "A", "Book", "Stock Opening", 20, 200, "O1", datetime(2026, 5, 1, 8)),
        make_tx("FS", "FS Team 07", "A", "Book", "Stock Used", 4, 40, "R1", datetime(2026, 5, 2, 9)),
        make_tx("KZN", "KZN Team 02", "B", "Pen", "Stock Opening", 10, 100, "O2", datetime(2026, 5, 1, 8)),
        make_tx("KZN", "KZN Team 02", "B", "Pen", "Stock Used", 3, 30, "R2", datetime(2026, 5, 3, 9)),
    ]
    monkeypatch.setattr("backend.app.api._current_transactions", lambda: txs)

    summary = _build_executive_summary(
        scope="team",
        regions=["FS", "KZN"],
        teams=["FS Team 07", "KZN Team 02"],
        item_codes=[],
        from_date=None,
        to_date=None,
    )

    assert summary["metrics"]["teams"] == 2
    assert summary["metrics"]["regions"] == 2
    assert summary["metrics"]["used_quantity"] == 7
    assert len(summary["usage_by_item"]) == 2
    assert len(summary["usage_by_region"]) == 2


def test_executive_report_exports_are_nonempty(monkeypatch):
    txs = [
        make_tx("FS", "FS Team 07", "A", "Book", "Stock Opening", 20, 200, "O1", datetime(2026, 5, 1, 8)),
        make_tx("FS", "FS Team 07", "A", "Book", "Stock Used", 4, 40, "R1", datetime(2026, 5, 2, 9)),
    ]
    monkeypatch.setattr("backend.app.api._current_transactions", lambda: txs)

    summary = _build_executive_summary(
        scope="team",
        regions=["FS"],
        teams=["FS Team 07"],
        item_codes=[],
        from_date=None,
        to_date=None,
    )

    text = build_text_report(summary)
    pdf = build_pdf_report(summary).getvalue()
    docx = build_docx_report(summary).getvalue()

    assert "STOCK INTELLIGENCE - EXECUTIVE REPORT" in text
    assert len(pdf) > 1000
    assert pdf.startswith(b"%PDF")
    assert len(docx) > 1000
    assert docx[:2] == b"PK"


def test_stock_table_matches_reference_columns_and_values(monkeypatch):
    txs = [
        make_tx("GAU", "Regional Office Gauteng", "CAN", "Can Cooler", "Stock Opening", 19, "288.61", "O1", datetime(2026, 9, 21, 8)),
        make_tx("GAU", "Regional Office Gauteng", "CAN", "Can Cooler", "Stock Used", 17, "258.23", "U1", datetime(2026, 9, 23, 9)),
    ]
    monkeypatch.setattr("backend.app.api._current_transactions", lambda: txs)
    summary = _build_executive_summary(scope="region", regions=["GAU"], teams=[], item_codes=[],
                                        from_date=date(2026, 9, 21), to_date=date(2026, 9, 28))
    assert summary["detail_columns"] == [
        "Date", "Operator", "Team", "Item", "21-Sept-26", "New Stock", "Stock Give Away",
        "28-Sept-26", "Variance", "Value per Item", "Opening Value", "New Stock Value",
        "SG Value", "Closing Value",
    ]
    row = summary["detail_rows"][0]
    assert (row["opening_quantity"], row["stock_give_away"], row["closing_quantity"], row["variance"]) == (19, 17, 2, 0)
    assert (row["value_per_item"], row["opening_value"], row["sg_value"], row["closing_value"]) == (15.19, 288.61, 258.23, 30.38)

    sheet = load_workbook(BytesIO(build_xlsx_table(summary).getvalue()), data_only=True).active
    assert [cell.value for cell in sheet[1]] == summary["detail_columns"]
    assert sheet["E2"].value == 19
    assert sheet["F2"].value is None
    assert sheet["I2"].value == 0
    assert sheet["N2"].value == 30.38
    assert build_pdf_table(summary).getvalue().startswith(b"%PDF")
    word = Document(BytesIO(build_docx_table(summary).getvalue()))
    assert [cell.text for cell in word.tables[0].rows[0].cells] == summary["detail_columns"]


def test_selection_meta_includes_every_region_team_and_item():
    txs = [
        make_tx("GAU", "GAU Team 11", "CAN", "Can Cooler", "Stock Opening", 1, 10, "O1", datetime(2026, 9, 21)),
        make_tx("GAU", "GAU Team 11", "HAT", "Hat", "Stock Opening", 1, 10, "O2", datetime(2026, 9, 21)),
        make_tx("FS", "FS Team 06", "PEN", "Pen", "Stock Opening", 1, 10, "O3", datetime(2026, 9, 21)),
    ]
    meta = _selection_meta(txs)
    assert set(meta["regions"]) == {tx.region for tx in txs}
    assert set(meta["teams"]) == {tx.team for tx in txs}
    assert set(meta["items"]) == {tx.item_code for tx in txs}
    assert meta["teams_by_region"] == {"FS": ["FS Team 06"], "GAU": ["GAU Team 11"]}
    assert meta["items_by_team"] == {"FS Team 06": ["PEN"], "GAU Team 11": ["CAN", "HAT"]}
