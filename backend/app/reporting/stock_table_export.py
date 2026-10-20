"""The shared stock movement table used by the executive exports."""

from __future__ import annotations

from collections import defaultdict
from datetime import date, datetime
from decimal import Decimal, ROUND_HALF_UP
from io import BytesIO
from xml.sax.saxutils import escape

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.shared import Inches, Pt, RGBColor
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A3, landscape
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Table, TableStyle

from backend.app.domain.models import StockTransaction
from backend.app.domain.transaction_rules import normalize_transaction_type

ZERO = Decimal("0")
CENT = Decimal("0.01")
HEADER_BLUE = "BFD7EB"
VARIANCE_GREEN = "C6EFCE"
GRID = "D5DEE6"


def _header_date(value: date) -> str:
    month = "Sept" if value.month == 9 else value.strftime("%b")
    return f"{value.day}-{month}-{value.strftime('%y')}"


def _money(value: Decimal | None) -> float | None:
    return float(value.quantize(CENT, rounding=ROUND_HALF_UP)) if value is not None else None


def _quantity(value: Decimal | None) -> float | None:
    return float(value) if value is not None else None


def build_stock_table(
    transactions: list[StockTransaction], from_date: date | None, to_date: date | None,
) -> tuple[list[str], list[dict[str, object]]]:
    dated = [tx.date.date() for tx in transactions if tx.date]
    start = from_date or (min(dated) if dated else date.today())
    end = to_date or (max(dated) if dated else start)
    columns = [
        "Date", "Operator", "Team", "Item", _header_date(start), "New Stock",
        "Stock Give Away", _header_date(end), "Variance", "Value per Item",
        "Opening Value", "New Stock Value", "SG Value", "Closing Value",
    ]
    grouped: dict[tuple[str, str, str], list[StockTransaction]] = defaultdict(list)
    for tx in transactions:
        grouped[(tx.region, tx.team, tx.item_code)].append(tx)

    rows: list[dict[str, object]] = []
    for (_, team, _), entries in sorted(grouped.items()):
        entries.sort(key=lambda tx: tx.date or datetime.max)
        openings = [tx for tx in entries if normalize_transaction_type(tx.transaction_type) == "Stock Opening"]
        baseline = [tx for tx in openings if tx.date and tx.date.date() <= start]
        first_opening = max(baseline, key=lambda tx: tx.date) if baseline else (openings[0] if openings else None)
        last_opening = openings[-1] if openings else None
        activity = [tx for tx in entries if tx.date and start <= tx.date.date() <= end]
        first_activity = activity[0] if activity else (first_opening or entries[0])
        new_stock = sum((tx.quantity for tx in activity if normalize_transaction_type(tx.transaction_type) in {"Intake", "Movement In"}), ZERO)
        stock_give_away = sum((tx.quantity for tx in activity if normalize_transaction_type(tx.transaction_type) in {"Stock Used", "Movement Out"}), ZERO)

        closing: Decimal | None = None
        if last_opening is not None:
            closing = last_opening.quantity
            for tx in entries:
                if tx.date and last_opening.date and tx.date > last_opening.date:
                    kind = normalize_transaction_type(tx.transaction_type)
                    if kind in {"Intake", "Movement In"}:
                        closing += tx.quantity
                    elif kind in {"Stock Used", "Movement Out"}:
                        closing -= tx.quantity

        opening = first_opening.quantity if first_opening else None
        variance = opening + new_stock - stock_give_away - closing if opening is not None and closing is not None else None
        priced = first_opening if first_opening and first_opening.quantity else next((tx for tx in activity if tx.quantity and tx.value), None)
        price = (priced.value / priced.quantity).quantize(CENT, rounding=ROUND_HALF_UP) if priced else None
        rows.append({
            "date": start.isoformat(),
            "operator": first_activity.operator,
            "team": team,
            "item": next((tx.item_name for tx in reversed(entries) if tx.item_name), entries[-1].item_code),
            "opening_quantity": _quantity(opening),
            "new_stock": _quantity(new_stock),
            "stock_give_away": _quantity(stock_give_away),
            "closing_quantity": _quantity(closing),
            "variance": _quantity(variance),
            "value_per_item": _money(price),
            "opening_value": _money(opening * price) if opening is not None and price is not None else None,
            "new_stock_value": _money(new_stock * price) if price is not None else None,
            "sg_value": _money(stock_give_away * price) if price is not None else None,
            "closing_value": _money(closing * price) if closing is not None and price is not None else None,
        })
    return columns, rows


def _row_values(row: dict[str, object]) -> list[object]:
    values = [row[key] for key in (
        "date", "operator", "team", "item", "opening_quantity", "new_stock",
        "stock_give_away", "closing_quantity", "variance", "value_per_item",
        "opening_value", "new_stock_value", "sg_value", "closing_value",
    )]
    if values[5] == 0:
        values[5] = None
    return values


def _formatted(value: object, index: int) -> str:
    if value is None:
        return ""
    if index == 0:
        return str(value).replace("-", "/")
    if index >= 4:
        number = f"{float(value):,.2f}"
        return number.replace(",", " ").replace(".", ",")
    return str(value)


def build_xlsx_table(summary: dict[str, object]) -> BytesIO:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Stock Report"
    headers = summary["detail_columns"]
    sheet.append(headers)
    widths = [14, 47, 38, 50, 15, 16, 19, 15, 13, 18, 19, 20, 16, 19]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.row_dimensions[1].height = 25
    thin = Side(style="thin", color=GRID)
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor=HEADER_BLUE)
        cell.font = Font(bold=True, color="12273A")
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        cell.border = Border(bottom=thin, right=thin)
    for row in summary["detail_rows"]:
        sheet.append(_row_values(row))
        row_number = sheet.max_row
        sheet.cell(row_number, 1).value = date.fromisoformat(row["date"])
        sheet.cell(row_number, 1).number_format = "yyyy/mm/dd"
        for index in (2, 3, 4):
            sheet.cell(row_number, index).data_type = "s"
        for index in range(5, 15):
            cell = sheet.cell(row_number, index)
            cell.number_format = "#,##0.00"
            cell.alignment = Alignment(horizontal="right")
        sheet.cell(row_number, 9).fill = PatternFill("solid", fgColor=VARIANCE_GREEN)
        for cell in sheet[row_number]:
            cell.border = Border(bottom=thin, right=thin)
    sheet.auto_filter.ref = f"A1:N{sheet.max_row}"
    sheet.freeze_panes = "E2"
    sheet.sheet_view.showGridLines = False
    sheet.print_options.horizontalCentered = True
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.print_title_rows = "1:1"
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output


def build_pdf_table(summary: dict[str, object]) -> BytesIO:
    output = BytesIO()
    doc = SimpleDocTemplate(output, pagesize=landscape(A3), leftMargin=8 * mm, rightMargin=8 * mm,
                            topMargin=10 * mm, bottomMargin=10 * mm, title="Stock Report")
    style = ParagraphStyle("Cell", fontName="Helvetica", fontSize=6.4, leading=8)
    bold = ParagraphStyle("Header", parent=style, fontName="Helvetica-Bold")
    matrix = [[Paragraph(escape(str(value)), bold) for value in summary["detail_columns"]]]
    for row in summary["detail_rows"]:
        matrix.append([Paragraph(escape(_formatted(value, index)), style) for index, value in enumerate(_row_values(row))])
    weights = [34, 97, 91, 120, 48, 49, 55, 48, 42, 53, 55, 57, 49, 55]
    usable_width = landscape(A3)[0] - 16 * mm
    scale = usable_width / sum(weights)
    table = Table(matrix, colWidths=[width * scale for width in weights], repeatRows=1, hAlign="LEFT")
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(f"#{HEADER_BLUE}")),
        ("BACKGROUND", (8, 1), (8, -1), colors.HexColor(f"#{VARIANCE_GREEN}")),
        ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor(f"#{GRID}")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F7FAFC")]),
    ]))
    doc.build([table])
    output.seek(0)
    return output


def build_docx_table(summary: dict[str, object]) -> BytesIO:
    document = Document()
    section = document.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Inches(16.54)
    section.page_height = Inches(11.69)
    section.left_margin = section.right_margin = Inches(0.35)
    section.top_margin = section.bottom_margin = Inches(0.4)
    table = document.add_table(rows=1, cols=14)
    table.style = "Table Grid"
    table.autofit = False
    weights = [0.68, 1.8, 1.75, 2.3, 0.82, 0.9, 1.05, 0.82, 0.8, 1.0, 1.05, 1.1, 0.9, 1.05]
    scale = 15.84 / sum(weights)
    for index, header in enumerate(summary["detail_columns"]):
        table.columns[index].width = Inches(weights[index] * scale)
        cell = table.rows[0].cells[index]
        cell.text = str(header)
        _shade_docx(cell, HEADER_BLUE)
        for run in cell.paragraphs[0].runs:
            run.bold = True
    table.rows[0]._tr.get_or_add_trPr().append(_repeat_table_header())
    for row in summary["detail_rows"]:
        cells = table.add_row().cells
        for index, value in enumerate(_row_values(row)):
            cells[index].text = _formatted(value, index)
            cells[index].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if index == 8:
                _shade_docx(cells[index], VARIANCE_GREEN)
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            cell.width = Inches(weights[index] * scale)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                for run in paragraph.runs:
                    run.font.name = "Aptos"
                    run.font.size = Pt(6.5)
                    run.font.color.rgb = RGBColor(18, 39, 58)
    output = BytesIO()
    document.save(output)
    output.seek(0)
    return output


def _shade_docx(cell, color: str) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    shade = OxmlElement("w:shd")
    shade.set(qn("w:fill"), color)
    cell._tc.get_or_add_tcPr().append(shade)


def _repeat_table_header():
    from docx.oxml import OxmlElement
    return OxmlElement("w:tblHeader")
