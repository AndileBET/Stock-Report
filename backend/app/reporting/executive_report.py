from __future__ import annotations

from datetime import datetime
from io import BytesIO
from typing import Any
from xml.sax.saxutils import escape

from docx import Document
from docx.enum.table import (
    WD_CELL_VERTICAL_ALIGNMENT,
)
from docx.enum.text import (
    WD_ALIGN_PARAGRAPH,
)
from docx.shared import (
    Inches,
    Pt,
    RGBColor,
)

from reportlab.lib import colors
from reportlab.lib.pagesizes import (
    A4,
    landscape,
)
from reportlab.lib.styles import (
    ParagraphStyle,
    getSampleStyleSheet,
)
from reportlab.lib.units import mm

from reportlab.platypus import (
    KeepTogether,
    PageBreak,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)


PURPLE = "5B2A86"
DEEP_PURPLE = "30144D"
MAGENTA = "A33B9F"

LAVENDER = "F3EDF9"
VERY_LIGHT_PURPLE = "FBF9FD"

INK = "24152F"
MUTED = "6D6078"

RISK = "B4234D"
GOOD = "197A55"

WHITE = "FFFFFF"


# =========================================================
# TEXT REPORT
# =========================================================

def build_text_report(
    summary: dict[str, Any],
) -> str:
    filters = summary["filters"]
    metrics = summary["metrics"]

    lines = [
        "STOCK INTELLIGENCE - EXECUTIVE REPORT",
        "=" * 72,
        "",
        f"Generated: {summary['generated_at']}",
        f"Scope: {summary['scope'].title()}",
        (
            "Period: "
            + _period_label(
                filters.get("from_date"),
                filters.get("to_date"),
            )
        ),
        "",
        "REPORT SCOPE",
        "-" * 72,
        "",
        "Provinces:",
        _format_text_selection(
            filters.get("regions")
        ),
        "",
        "Teams:",
        _format_text_selection(
            filters.get("teams")
        ),
        "",
        "Items:",
        _format_text_selection(
            filters.get("item_codes")
        ),
        "",
        "EXECUTIVE METRICS",
        "-" * 72,
        (
            f"Expected stock: "
            f"{metrics['expected_quantity']:,.0f}"
        ),
        (
            f"Stock used: "
            f"{metrics['used_quantity']:,.0f}"
        ),
        (
            f"Usage value: "
            f"{metrics['usage_value']:,.2f}"
        ),
        (
            f"Closing stock: "
            f"{metrics['closing_quantity']:,.0f}"
        ),
        (
            f"Shortage quantity: "
            f"{metrics['shortage_quantity']:,.0f}"
        ),
        (
            f"Shortage value: "
            f"{metrics['shortage_value']:,.2f}"
        ),
        (
            f"Provinces represented: "
            f"{metrics['regions']}"
        ),
        (
            f"Teams represented: "
            f"{metrics['teams']}"
        ),
        (
            f"Items represented: "
            f"{metrics['items']}"
        ),
        "",
        "USAGE BY ITEM",
        "-" * 72,
    ]

    for row in summary["usage_by_item"]:
        lines.append(
            f"{row['region']} | "
            f"{row['team']} | "
            f"{row['item_code']} | "
            f"{row['item_name']} | "
            f"Used {row['used_quantity']:,.0f} | "
            f"Value {row['usage_value']:,.2f}"
        )

    shortage_rows = [
        row
        for row
        in summary["variance_records"]
        if row.get(
            "shortage_quantity",
            0,
        ) > 0
    ]

    lines.extend(
        [
            "",
            "SHORTAGE EXPOSURE",
            "-" * 72,
        ]
    )

    if not shortage_rows:
        lines.append(
            "No shortage exposure identified "
            "for the selected scope."
        )

    for row in shortage_rows[:50]:
        lines.append(
            f"{row['region']} | "
            f"{row['team']} | "
            f"{row['item_code']} | "
            f"{row['item_name']} | "
            f"Shortage "
            f"{row['shortage_quantity']:,.0f} | "
            f"Value "
            f"{row['shortage_value']:,.2f}"
        )

    lines.extend(
        [
            "",
            "REPORTING NOTE",
            "-" * 72,
            (
                "Opening Stock is treated as a daily "
                "starting snapshot, not an inflow. "
                "Count transactions are excluded from "
                "stock calculations. Expected stock is "
                "Opening + Intake + Movement In - Movement Out. "
                "Closing stock is Expected Stock - Stock Used."
            ),
        ]
    )

    return "\n".join(
        lines
    )


# =========================================================
# WORD / DOCX
# =========================================================

def build_docx_report(
    summary: dict[str, Any],
) -> BytesIO:
    doc = Document()

    section = doc.sections[0]

    section.top_margin = Inches(
        0.55
    )
    section.bottom_margin = Inches(
        0.55
    )
    section.left_margin = Inches(
        0.62
    )
    section.right_margin = Inches(
        0.62
    )

    styles = doc.styles

    styles["Normal"].font.name = (
        "Aptos"
    )

    styles["Normal"].font.size = (
        Pt(9.5)
    )

    # -----------------------------------------------------
    # TITLE
    # -----------------------------------------------------

    title = doc.add_paragraph()

    title.alignment = (
        WD_ALIGN_PARAGRAPH.LEFT
    )

    run = title.add_run(
        "STOCK INTELLIGENCE"
    )

    run.bold = True
    run.font.size = Pt(11)
    run.font.color.rgb = (
        RGBColor.from_string(
            PURPLE
        )
    )

    heading = doc.add_paragraph()

    heading.paragraph_format.space_after = (
        Pt(4)
    )

    run = heading.add_run(
        "Executive Report"
    )

    run.bold = True
    run.font.size = Pt(25)
    run.font.color.rgb = (
        RGBColor.from_string(
            DEEP_PURPLE
        )
    )

    filters = summary[
        "filters"
    ]

    meta = doc.add_paragraph()

    meta.paragraph_format.space_after = (
        Pt(14)
    )

    r = meta.add_run(
        "Scope: "
    )

    r.bold = True

    meta.add_run(
        summary["scope"].title()
    )

    meta.add_run(
        "    •    "
    )

    r = meta.add_run(
        "Period: "
    )

    r.bold = True

    meta.add_run(
        _period_label(
            filters.get(
                "from_date"
            ),
            filters.get(
                "to_date"
            ),
        )
    )

    # -----------------------------------------------------
    # REPORT SCOPE
    # -----------------------------------------------------

    doc.add_heading(
        "Report scope",
        level=1,
    )

    scope_table = doc.add_table(
        rows=0,
        cols=2,
    )

    scope_table.style = (
        "Table Grid"
    )

    _add_docx_scope_row(
        scope_table,
        "Provinces",
        filters.get(
            "regions"
        ),
    )

    _add_docx_scope_row(
        scope_table,
        "Teams",
        filters.get(
            "teams"
        ),
    )

    _add_docx_scope_row(
        scope_table,
        "Items",
        filters.get(
            "item_codes"
        ),
    )

    _add_docx_scope_text_row(
        scope_table,
        "Generated",
        _format_generated_timestamp(
            summary[
                "generated_at"
            ]
        ),
    )

    doc.add_paragraph()

    # -----------------------------------------------------
    # EXECUTIVE OVERVIEW
    # -----------------------------------------------------

    doc.add_heading(
        "Executive overview",
        level=1,
    )

    metrics = summary[
        "metrics"
    ]

    metric_table = doc.add_table(
        rows=2,
        cols=6,
    )

    metric_table.style = (
        "Table Grid"
    )

    headers = [
        "Expected Stock",
        "Stock Used",
        "Usage Value",
        "Closing Stock",
        "Shortage",
        "Shortage Value",
    ]

    values = [
        (
            f"{metrics['expected_quantity']:,.0f}"
        ),
        (
            f"{metrics['used_quantity']:,.0f}"
        ),
        (
            f"{metrics['usage_value']:,.2f}"
        ),
        (
            f"{metrics['closing_quantity']:,.0f}"
        ),
        (
            f"{metrics['shortage_quantity']:,.0f}"
        ),
        (
            f"{metrics['shortage_value']:,.2f}"
        ),
    ]

    for index, text in enumerate(
        headers
    ):
        metric_table.cell(
            0,
            index,
        ).text = text

        metric_table.cell(
            1,
            index,
        ).text = values[
            index
        ]

        _shade_docx_cell(
            metric_table.cell(
                0,
                index,
            ),
            DEEP_PURPLE,
        )

        header_para = (
            metric_table
            .cell(
                0,
                index,
            )
            .paragraphs[
                0
            ]
        )

        header_para.alignment = (
            WD_ALIGN_PARAGRAPH.CENTER
        )

        for header_run in (
            header_para.runs
        ):
            header_run.bold = (
                True
            )

            header_run.font.color.rgb = (
                RGBColor.from_string(
                    WHITE
                )
            )

        value_para = (
            metric_table
            .cell(
                1,
                index,
            )
            .paragraphs[
                0
            ]
        )

        value_para.alignment = (
            WD_ALIGN_PARAGRAPH.CENTER
        )

        for value_run in (
            value_para.runs
        ):
            value_run.bold = True

    doc.add_paragraph()

    # -----------------------------------------------------
    # USAGE TABLE
    # -----------------------------------------------------

    doc.add_heading(
        "Usage by item",
        level=1,
    )

    usage_table = doc.add_table(
        rows=1,
        cols=6,
    )

    usage_table.style = (
        "Table Grid"
    )

    usage_headers = [
        "Region",
        "Team",
        "Item",
        "Code",
        "Used",
        "Value",
    ]

    for index, text in enumerate(
        usage_headers
    ):
        cell = usage_table.cell(
            0,
            index,
        )

        cell.text = text

        _shade_docx_cell(
            cell,
            PURPLE,
        )

        for header_run in (
            cell
            .paragraphs[
                0
            ]
            .runs
        ):
            header_run.bold = (
                True
            )

            header_run.font.color.rgb = (
                RGBColor.from_string(
                    WHITE
                )
            )

    for row in summary[
        "usage_by_item"
    ]:
        cells = (
            usage_table
            .add_row()
            .cells
        )

        values = [
            row["region"],
            row["team"],
            row["item_name"],
            row["item_code"],
            (
                f"{row['used_quantity']:,.0f}"
            ),
            (
                f"{row['usage_value']:,.2f}"
            ),
        ]

        for index, value in enumerate(
            values
        ):
            cells[
                index
            ].text = str(
                value
            )

            cells[
                index
            ].vertical_alignment = (
                WD_CELL_VERTICAL_ALIGNMENT.CENTER
            )

            if index >= 4:
                cells[
                    index
                ].paragraphs[
                    0
                ].alignment = (
                    WD_ALIGN_PARAGRAPH.RIGHT
                )

    # -----------------------------------------------------
    # SHORTAGES
    # -----------------------------------------------------

    shortage_rows = [
        row
        for row
        in summary[
            "variance_records"
        ]
        if row.get(
            "shortage_quantity",
            0,
        ) > 0
    ]

    if shortage_rows:
        doc.add_paragraph()

        doc.add_heading(
            "Shortage exposure",
            level=1,
        )

        shortage_table = (
            doc.add_table(
                rows=1,
                cols=5,
            )
        )

        shortage_table.style = (
            "Table Grid"
        )

        shortage_headers = [
            "Region",
            "Team",
            "Item",
            "Shortage",
            "Value",
        ]

        for index, text in enumerate(
            shortage_headers
        ):
            cell = (
                shortage_table
                .cell(
                    0,
                    index,
                )
            )

            cell.text = text

            _shade_docx_cell(
                cell,
                MAGENTA,
            )

            for header_run in (
                cell
                .paragraphs[
                    0
                ]
                .runs
            ):
                header_run.bold = (
                    True
                )

                header_run.font.color.rgb = (
                    RGBColor.from_string(
                        WHITE
                    )
                )

        for row in shortage_rows[
            :50
        ]:
            cells = (
                shortage_table
                .add_row()
                .cells
            )

            values = [
                row["region"],
                row["team"],
                row["item_name"],
                (
                    f"{row['shortage_quantity']:,.0f}"
                ),
                (
                    f"{row['shortage_value']:,.2f}"
                ),
            ]

            for index, value in enumerate(
                values
            ):
                cells[
                    index
                ].text = str(
                    value
                )

                if index >= 3:
                    cells[
                        index
                    ].paragraphs[
                        0
                    ].alignment = (
                        WD_ALIGN_PARAGRAPH.RIGHT
                    )

    # -----------------------------------------------------
    # NOTE
    # -----------------------------------------------------

    doc.add_paragraph()

    note = doc.add_paragraph()

    note_run = note.add_run(
        "Reporting note: "
    )

    note_run.bold = True
    note_run.font.color.rgb = (
        RGBColor.from_string(
            PURPLE
        )
    )

    note.add_run(
        "Opening Stock is a daily starting snapshot rather "
        "than an additive inflow. Count transactions are "
        "excluded. Expected stock equals Opening + Intake + "
        "Movement In - Movement Out, and closing stock equals "
        "Expected Stock - Stock Used."
    )

    buffer = BytesIO()

    doc.save(
        buffer
    )

    buffer.seek(0)

    return buffer


# =========================================================
# PDF
# =========================================================

def build_pdf_report(
    summary: dict[str, Any],
) -> BytesIO:
    buffer = BytesIO()

    doc = SimpleDocTemplate(
        buffer,
        pagesize=landscape(
            A4
        ),
        leftMargin=13 * mm,
        rightMargin=13 * mm,
        topMargin=14 * mm,
        bottomMargin=15 * mm,
        title=(
            "Stock Intelligence Executive Report"
        ),
        author=(
            "Stock Intelligence"
        ),
    )

    styles = (
        getSampleStyleSheet()
    )

    styles.add(
        ParagraphStyle(
            name="CoverTitle",
            parent=styles[
                "Title"
            ],
            fontName=(
                "Helvetica-Bold"
            ),
            fontSize=23,
            leading=27,
            textColor=(
                colors.HexColor(
                    f"#{DEEP_PURPLE}"
                )
            ),
            spaceAfter=6,
        )
    )

    styles.add(
        ParagraphStyle(
            name="Eyebrow",
            parent=styles[
                "Normal"
            ],
            fontName=(
                "Helvetica-Bold"
            ),
            fontSize=8,
            leading=10,
            textColor=(
                colors.HexColor(
                    f"#{PURPLE}"
                )
            ),
            spaceAfter=3,
        )
    )

    styles.add(
        ParagraphStyle(
            name="Section",
            parent=styles[
                "Heading2"
            ],
            fontName=(
                "Helvetica-Bold"
            ),
            fontSize=13,
            leading=16,
            textColor=(
                colors.HexColor(
                    f"#{DEEP_PURPLE}"
                )
            ),
            spaceBefore=9,
            spaceAfter=6,
        )
    )

    styles.add(
        ParagraphStyle(
            name="BodySmall",
            parent=styles[
                "BodyText"
            ],
            fontName=(
                "Helvetica"
            ),
            fontSize=8.5,
            leading=11.5,
            textColor=(
                colors.HexColor(
                    f"#{INK}"
                )
            ),
        )
    )

    styles.add(
        ParagraphStyle(
            name="ScopeLabel",
            parent=styles[
                "BodyText"
            ],
            fontName=(
                "Helvetica-Bold"
            ),
            fontSize=8,
            leading=10,
            textColor=(
                colors.HexColor(
                    f"#{PURPLE}"
                )
            ),
        )
    )

    styles.add(
        ParagraphStyle(
            name="ScopeValue",
            parent=styles[
                "BodyText"
            ],
            fontName=(
                "Helvetica"
            ),
            fontSize=8,
            leading=11,
            textColor=(
                colors.HexColor(
                    f"#{INK}"
                )
            ),
        )
    )

    styles.add(
        ParagraphStyle(
            name="TableCell",
            parent=styles[
                "BodyText"
            ],
            fontName=(
                "Helvetica"
            ),
            fontSize=7.2,
            leading=9,
            textColor=(
                colors.HexColor(
                    f"#{INK}"
                )
            ),
        )
    )

    styles.add(
        ParagraphStyle(
            name="TableHeader",
            parent=styles[
                "BodyText"
            ],
            fontName=(
                "Helvetica-Bold"
            ),
            fontSize=7.3,
            leading=9,
            textColor=(
                colors.white
            ),
        )
    )

    story: list[Any] = []

    # -----------------------------------------------------
    # TITLE
    # -----------------------------------------------------

    story.append(
        Paragraph(
            "STOCK INTELLIGENCE",
            styles[
                "Eyebrow"
            ],
        )
    )

    story.append(
        Paragraph(
            "Executive Report",
            styles[
                "CoverTitle"
            ],
        )
    )

    filters = summary[
        "filters"
    ]

    story.append(
        Paragraph(
            (
                f"<b>Scope:</b> "
                f"{escape(summary['scope'].title())}"
                f"&nbsp;&nbsp;&nbsp;&nbsp;"
                f"<b>Period:</b> "
                f"{escape(_period_label(filters.get('from_date'), filters.get('to_date')))}"
            ),
            styles[
                "BodySmall"
            ],
        )
    )

    story.append(
        Spacer(
            1,
            4 * mm,
        )
    )

    # -----------------------------------------------------
    # REPORT SCOPE
    # -----------------------------------------------------

    story.append(
        Paragraph(
            "Report scope",
            styles[
                "Section"
            ],
        )
    )

    story.extend(
        _build_pdf_scope_rows(
            "Provinces",
            filters.get(
                "regions"
            ),
            styles,
        )
    )

    story.extend(
        _build_pdf_scope_rows(
            "Teams",
            filters.get(
                "teams"
            ),
            styles,
        )
    )

    story.extend(
        _build_pdf_scope_rows(
            "Items",
            filters.get(
                "item_codes"
            ),
            styles,
        )
    )

    story.extend(
        _build_pdf_scope_rows(
            "Generated",
            [
                _format_generated_timestamp(
                    summary[
                        "generated_at"
                    ]
                )
            ],
            styles,
        )
    )

    story.append(
        Spacer(
            1,
            4 * mm,
        )
    )

    # -----------------------------------------------------
    # EXECUTIVE OVERVIEW
    # -----------------------------------------------------

    story.append(
        Paragraph(
            "Executive overview",
            styles[
                "Section"
            ],
        )
    )

    metrics = summary[
        "metrics"
    ]

    metric_data = [
        [
            Paragraph(
                "Expected Stock",
                styles[
                    "TableHeader"
                ],
            ),
            Paragraph(
                "Stock Used",
                styles[
                    "TableHeader"
                ],
            ),
            Paragraph(
                "Usage Value",
                styles[
                    "TableHeader"
                ],
            ),
            Paragraph(
                "Closing Stock",
                styles[
                    "TableHeader"
                ],
            ),
            Paragraph(
                "Shortage",
                styles[
                    "TableHeader"
                ],
            ),
            Paragraph(
                "Shortage Value",
                styles[
                    "TableHeader"
                ],
            ),
        ],
        [
            (
                f"{metrics['expected_quantity']:,.0f}"
            ),
            (
                f"{metrics['used_quantity']:,.0f}"
            ),
            (
                f"{metrics['usage_value']:,.2f}"
            ),
            (
                f"{metrics['closing_quantity']:,.0f}"
            ),
            (
                f"{metrics['shortage_quantity']:,.0f}"
            ),
            (
                f"{metrics['shortage_value']:,.2f}"
            ),
        ],
    ]

    metric_table = Table(
        metric_data,
        colWidths=[
            44 * mm
        ] * 6,
    )

    metric_table.setStyle(
        TableStyle(
            [
                (
                    "BACKGROUND",
                    (
                        0,
                        0,
                    ),
                    (
                        -1,
                        0,
                    ),
                    colors.HexColor(
                        f"#{DEEP_PURPLE}"
                    ),
                ),
                (
                    "TEXTCOLOR",
                    (
                        0,
                        0,
                    ),
                    (
                        -1,
                        0,
                    ),
                    colors.white,
                ),
                (
                    "FONTNAME",
                    (
                        0,
                        1,
                    ),
                    (
                        -1,
                        1,
                    ),
                    "Helvetica-Bold",
                ),
                (
                    "FONTSIZE",
                    (
                        0,
                        1,
                    ),
                    (
                        -1,
                        1,
                    ),
                    9,
                ),
                (
                    "ALIGN",
                    (
                        0,
                        0,
                    ),
                    (
                        -1,
                        -1,
                    ),
                    "CENTER",
                ),
                (
                    "VALIGN",
                    (
                        0,
                        0,
                    ),
                    (
                        -1,
                        -1,
                    ),
                    "MIDDLE",
                ),
                (
                    "BACKGROUND",
                    (
                        0,
                        1,
                    ),
                    (
                        -1,
                        1,
                    ),
                    colors.HexColor(
                        "#FBF9FD"
                    ),
                ),
                (
                    "BOX",
                    (
                        0,
                        0,
                    ),
                    (
                        -1,
                        -1,
                    ),
                    0.6,
                    colors.HexColor(
                        "#D9CDE6"
                    ),
                ),
                (
                    "INNERGRID",
                    (
                        0,
                        0,
                    ),
                    (
                        -1,
                        -1,
                    ),
                    0.3,
                    colors.HexColor(
                        "#E8E0EE"
                    ),
                ),
                (
                    "TOPPADDING",
                    (
                        0,
                        0,
                    ),
                    (
                        -1,
                        -1,
                    ),
                    7,
                ),
                (
                    "BOTTOMPADDING",
                    (
                        0,
                        0,
                    ),
                    (
                        -1,
                        -1,
                    ),
                    7,
                ),
            ]
        )
    )

    story.append(
        metric_table
    )

    story.append(
        Spacer(
            1,
            4 * mm,
        )
    )

    # -----------------------------------------------------
    # USAGE
    # -----------------------------------------------------

    story.append(
        Paragraph(
            "Usage by item",
            styles[
                "Section"
            ],
        )
    )

    usage_rows = [
        [
            Paragraph(
                "Region",
                styles[
                    "TableHeader"
                ],
            ),
            Paragraph(
                "Team",
                styles[
                    "TableHeader"
                ],
            ),
            Paragraph(
                "Item",
                styles[
                    "TableHeader"
                ],
            ),
            Paragraph(
                "Code",
                styles[
                    "TableHeader"
                ],
            ),
            Paragraph(
                "Used",
                styles[
                    "TableHeader"
                ],
            ),
            Paragraph(
                "Usage Value",
                styles[
                    "TableHeader"
                ],
            ),
        ]
    ]

    for row in summary[
        "usage_by_item"
    ]:
        usage_rows.append(
            [
                Paragraph(
                    escape(
                        str(
                            row[
                                "region"
                            ]
                        )
                    ),
                    styles[
                        "TableCell"
                    ],
                ),
                Paragraph(
                    escape(
                        str(
                            row[
                                "team"
                            ]
                        )
                    ),
                    styles[
                        "TableCell"
                    ],
                ),
                Paragraph(
                    escape(
                        str(
                            row[
                                "item_name"
                            ]
                        )
                    ),
                    styles[
                        "TableCell"
                    ],
                ),
                Paragraph(
                    escape(
                        str(
                            row[
                                "item_code"
                            ]
                        )
                    ),
                    styles[
                        "TableCell"
                    ],
                ),
                (
                    f"{row['used_quantity']:,.0f}"
                ),
                (
                    f"{row['usage_value']:,.2f}"
                ),
            ]
        )

    usage_table = Table(
        usage_rows,
        colWidths=[
            25 * mm,
            48 * mm,
            72 * mm,
            34 * mm,
            24 * mm,
            31 * mm,
        ],
        repeatRows=1,
    )

    usage_table.setStyle(
        TableStyle(
            [
                (
                    "BACKGROUND",
                    (
                        0,
                        0,
                    ),
                    (
                        -1,
                        0,
                    ),
                    colors.HexColor(
                        f"#{PURPLE}"
                    ),
                ),
                (
                    "TEXTCOLOR",
                    (
                        0,
                        0,
                    ),
                    (
                        -1,
                        0,
                    ),
                    colors.white,
                ),
                (
                    "GRID",
                    (
                        0,
                        0,
                    ),
                    (
                        -1,
                        -1,
                    ),
                    0.25,
                    colors.HexColor(
                        "#E6DDEB"
                    ),
                ),
                (
                    "ROWBACKGROUNDS",
                    (
                        0,
                        1,
                    ),
                    (
                        -1,
                        -1,
                    ),
                    [
                        colors.white,
                        colors.HexColor(
                            "#FCFAFD"
                        ),
                    ],
                ),
                (
                    "ALIGN",
                    (
                        4,
                        1,
                    ),
                    (
                        -1,
                        -1,
                    ),
                    "RIGHT",
                ),
                (
                    "VALIGN",
                    (
                        0,
                        0,
                    ),
                    (
                        -1,
                        -1,
                    ),
                    "TOP",
                ),
                (
                    "TOPPADDING",
                    (
                        0,
                        0,
                    ),
                    (
                        -1,
                        -1,
                    ),
                    4,
                ),
                (
                    "BOTTOMPADDING",
                    (
                        0,
                        0,
                    ),
                    (
                        -1,
                        -1,
                    ),
                    4,
                ),
                (
                    "LEFTPADDING",
                    (
                        0,
                        0,
                    ),
                    (
                        -1,
                        -1,
                    ),
                    5,
                ),
                (
                    "RIGHTPADDING",
                    (
                        0,
                        0,
                    ),
                    (
                        -1,
                        -1,
                    ),
                    5,
                ),
            ]
        )
    )

    story.append(
        usage_table
    )

    # -----------------------------------------------------
    # SHORTAGE PAGE
    # -----------------------------------------------------

    shortage_rows = [
        row
        for row
        in summary[
            "variance_records"
        ]
        if row.get(
            "shortage_quantity",
            0,
        ) > 0
    ]

    if shortage_rows:
        story.append(
            PageBreak()
        )

        story.append(
            Paragraph(
                "Shortage exposure",
                styles[
                    "Section"
                ],
            )
        )

        rows = [
            [
                Paragraph(
                    "Region",
                    styles[
                        "TableHeader"
                    ],
                ),
                Paragraph(
                    "Team",
                    styles[
                        "TableHeader"
                    ],
                ),
                Paragraph(
                    "Item",
                    styles[
                        "TableHeader"
                    ],
                ),
                Paragraph(
                    "Shortage",
                    styles[
                        "TableHeader"
                    ],
                ),
                Paragraph(
                    "Value",
                    styles[
                        "TableHeader"
                    ],
                ),
                Paragraph(
                    "Status",
                    styles[
                        "TableHeader"
                    ],
                ),
            ]
        ]

        for row in shortage_rows[
            :60
        ]:
            rows.append(
                [
                    Paragraph(
                        escape(
                            str(
                                row[
                                    "region"
                                ]
                            )
                        ),
                        styles[
                            "TableCell"
                        ],
                    ),
                    Paragraph(
                        escape(
                            str(
                                row[
                                    "team"
                                ]
                            )
                        ),
                        styles[
                            "TableCell"
                        ],
                    ),
                    Paragraph(
                        escape(
                            str(
                                row[
                                    "item_name"
                                ]
                            )
                        ),
                        styles[
                            "TableCell"
                        ],
                    ),
                    (
                        f"{row['shortage_quantity']:,.0f}"
                    ),
                    (
                        f"{row['shortage_value']:,.2f}"
                    ),
                    Paragraph(
                        escape(
                            str(
                                row.get(
                                    "variance_status",
                                    "Shortage",
                                )
                            )
                        ),
                        styles[
                            "TableCell"
                        ],
                    ),
                ]
            )

        shortage_table = Table(
            rows,
            colWidths=[
                27 * mm,
                48 * mm,
                75 * mm,
                28 * mm,
                31 * mm,
                28 * mm,
            ],
            repeatRows=1,
        )

        shortage_table.setStyle(
            TableStyle(
                [
                    (
                        "BACKGROUND",
                        (
                            0,
                            0,
                        ),
                        (
                            -1,
                            0,
                        ),
                        colors.HexColor(
                            f"#{MAGENTA}"
                        ),
                    ),
                    (
                        "TEXTCOLOR",
                        (
                            0,
                            0,
                        ),
                        (
                            -1,
                            0,
                        ),
                        colors.white,
                    ),
                    (
                        "GRID",
                        (
                            0,
                            0,
                        ),
                        (
                            -1,
                            -1,
                        ),
                        0.25,
                        colors.HexColor(
                            "#E8D5DF"
                        ),
                    ),
                    (
                        "ROWBACKGROUNDS",
                        (
                            0,
                            1,
                        ),
                        (
                            -1,
                            -1,
                        ),
                        [
                            colors.white,
                            colors.HexColor(
                                "#FFF8FA"
                            ),
                        ],
                    ),
                    (
                        "TEXTCOLOR",
                        (
                            5,
                            1,
                        ),
                        (
                            5,
                            -1,
                        ),
                        colors.HexColor(
                            f"#{RISK}"
                        ),
                    ),
                    (
                        "ALIGN",
                        (
                            3,
                            1,
                        ),
                        (
                            4,
                            -1,
                        ),
                        "RIGHT",
                    ),
                    (
                        "VALIGN",
                        (
                            0,
                            0,
                        ),
                        (
                            -1,
                            -1,
                        ),
                        "TOP",
                    ),
                    (
                        "TOPPADDING",
                        (
                            0,
                            0,
                        ),
                        (
                            -1,
                            -1,
                        ),
                        4,
                    ),
                    (
                        "BOTTOMPADDING",
                        (
                            0,
                            0,
                        ),
                        (
                            -1,
                            -1,
                        ),
                        4,
                    ),
                ]
            )
        )

        story.append(
            shortage_table
        )

    # -----------------------------------------------------
    # NOTE
    # -----------------------------------------------------

    story.append(
        Spacer(
            1,
            5 * mm,
        )
    )

    story.append(
        Paragraph(
            "Reporting note",
            styles[
                "Section"
            ],
        )
    )

    story.append(
        Paragraph(
            (
                "Opening Stock is a daily starting snapshot "
                "rather than an additive inflow. "
                "Count transactions are excluded. "
                "Expected stock equals Opening + Intake + "
                "Movement In - Movement Out, and closing stock "
                "equals Expected Stock - Stock Used."
            ),
            styles[
                "BodySmall"
            ],
        )
    )

    doc.build(
        story,
        onFirstPage=
            _pdf_footer,
        onLaterPages=
            _pdf_footer,
    )

    buffer.seek(0)

    return buffer


# =========================================================
# PDF SCOPE HELPERS
# =========================================================

def _build_pdf_scope_rows(
    label: str,
    values: list[str] | None,
    styles,
) -> list[Any]:
    """
    Build a scope block that can safely span pages.

    Values are chunked so one enormous Table row never becomes
    taller than the available page.
    """

    cleaned = [
        str(value).strip()
        for value in (
            values or []
        )
        if str(
            value
        ).strip()
    ]

    if not cleaned:
        cleaned = [
            "All"
        ]

    chunks = _chunk_values(
        cleaned,
        size=8,
    )

    rows = []

    for index, chunk in enumerate(
        chunks
    ):
        display = (
            " • ".join(
                escape(
                    value
                )
                for value
                in chunk
            )
        )

        if (
            index == 0
            and cleaned !=
            ["All"]
        ):
            display = (
                f"<b>{len(cleaned)} selected</b>"
                f"<br/>{display}"
            )

        rows.append(
            [
                Paragraph(
                    (
                        escape(
                            label
                        )
                        if index == 0
                        else ""
                    ),
                    styles[
                        "ScopeLabel"
                    ],
                ),
                Paragraph(
                    display,
                    styles[
                        "ScopeValue"
                    ],
                ),
            ]
        )

    table = Table(
        rows,
        colWidths=[
            31 * mm,
            225 * mm,
        ],
        hAlign="LEFT",
    )

    table.setStyle(
        TableStyle(
            [
                (
                    "BACKGROUND",
                    (
                        0,
                        0,
                    ),
                    (
                        0,
                        -1,
                    ),
                    colors.HexColor(
                        f"#{LAVENDER}"
                    ),
                ),
                (
                    "BACKGROUND",
                    (
                        1,
                        0,
                    ),
                    (
                        1,
                        -1,
                    ),
                    colors.HexColor(
                        "#FFFFFF"
                    ),
                ),
                (
                    "BOX",
                    (
                        0,
                        0,
                    ),
                    (
                        -1,
                        -1,
                    ),
                    0.4,
                    colors.HexColor(
                        "#DFD4E8"
                    ),
                ),
                (
                    "INNERGRID",
                    (
                        0,
                        0,
                    ),
                    (
                        -1,
                        -1,
                    ),
                    0.2,
                    colors.HexColor(
                        "#EEE7F3"
                    ),
                ),
                (
                    "VALIGN",
                    (
                        0,
                        0,
                    ),
                    (
                        -1,
                        -1,
                    ),
                    "TOP",
                ),
                (
                    "LEFTPADDING",
                    (
                        0,
                        0,
                    ),
                    (
                        -1,
                        -1,
                    ),
                    7,
                ),
                (
                    "RIGHTPADDING",
                    (
                        0,
                        0,
                    ),
                    (
                        -1,
                        -1,
                    ),
                    7,
                ),
                (
                    "TOPPADDING",
                    (
                        0,
                        0,
                    ),
                    (
                        -1,
                        -1,
                    ),
                    6,
                ),
                (
                    "BOTTOMPADDING",
                    (
                        0,
                        0,
                    ),
                    (
                        -1,
                        -1,
                    ),
                    6,
                ),
            ]
        )
    )

    return [
        table,
        Spacer(
            1,
            2 * mm,
        ),
    ]


# =========================================================
# WORD HELPERS
# =========================================================

def _add_docx_scope_row(
    table,
    label: str,
    values: list[str] | None,
) -> None:
    row = table.add_row()

    label_cell = (
        row.cells[
            0
        ]
    )

    value_cell = (
        row.cells[
            1
        ]
    )

    label_cell.text = (
        label
    )

    _shade_docx_cell(
        label_cell,
        LAVENDER,
    )

    for run in (
        label_cell
        .paragraphs[
            0
        ]
        .runs
    ):
        run.bold = True

        run.font.color.rgb = (
            RGBColor.from_string(
                PURPLE
            )
        )

    cleaned = [
        str(value).strip()
        for value in (
            values or []
        )
        if str(
            value
        ).strip()
    ]

    if not cleaned:
        cleaned = [
            "All"
        ]

    # Remove initial empty paragraph.
    value_cell.text = ""

    count_para = (
        value_cell
        .paragraphs[
            0
        ]
    )

    if cleaned != [
        "All"
    ]:
        run = (
            count_para
            .add_run(
                f"{len(cleaned)} selected"
            )
        )

        run.bold = True

    else:
        count_para.add_run(
            "All"
        )

    if cleaned != [
        "All"
    ]:
        for chunk in _chunk_values(
            cleaned,
            size=8,
        ):
            paragraph = (
                value_cell
                .add_paragraph()
            )

            paragraph.add_run(
                " • ".join(
                    chunk
                )
            )

    for cell in (
        row.cells
    ):
        cell.vertical_alignment = (
            WD_CELL_VERTICAL_ALIGNMENT.TOP
        )


def _add_docx_scope_text_row(
    table,
    label: str,
    value: str,
) -> None:
    row = table.add_row()

    row.cells[
        0
    ].text = label

    row.cells[
        1
    ].text = value

    _shade_docx_cell(
        row.cells[
            0
        ],
        LAVENDER,
    )

    for run in (
        row.cells[
            0
        ]
        .paragraphs[
            0
        ]
        .runs
    ):
        run.bold = True

        run.font.color.rgb = (
            RGBColor.from_string(
                PURPLE
            )
        )


def _shade_docx_cell(
    cell,
    hex_color: str,
) -> None:
    from docx.oxml import (
        OxmlElement,
    )

    from docx.oxml.ns import (
        qn,
    )

    tc_pr = (
        cell
        ._tc
        .get_or_add_tcPr()
    )

    shd = tc_pr.find(
        qn(
            "w:shd"
        )
    )

    if shd is None:
        shd = (
            OxmlElement(
                "w:shd"
            )
        )

        tc_pr.append(
            shd
        )

    shd.set(
        qn(
            "w:fill"
        ),
        hex_color,
    )


# =========================================================
# GENERAL HELPERS
# =========================================================

def _chunk_values(
    values: list[str],
    size: int = 8,
) -> list[list[str]]:
    return [
        values[
            index:
            index + size
        ]
        for index
        in range(
            0,
            len(values),
            size,
        )
    ]


def _format_text_selection(
    values: list[str] | None,
) -> str:
    if not values:
        return "  All"

    return "\n".join(
        (
            f"  • {value}"
            for value
            in values
        )
    )


def _join_or_all(
    values: list[str] | None,
) -> str:
    if not values:
        return "All"

    return ", ".join(
        values
    )


def _period_label(
    from_date: str | None,
    to_date: str | None,
) -> str:
    if (
        from_date
        and to_date
    ):
        return (
            f"{from_date} to "
            f"{to_date}"
        )

    if from_date:
        return (
            f"From {from_date}"
        )

    if to_date:
        return (
            f"Through {to_date}"
        )

    return (
        "All available data"
    )


def _format_generated_timestamp(
    value: str,
) -> str:
    try:
        generated = (
            datetime.fromisoformat(
                value
            )
        )

        return generated.strftime(
            "%d %B %Y at %H:%M"
        )

    except (
        ValueError,
        TypeError,
    ):
        return str(
            value
        )


# =========================================================
# PDF FOOTER
# =========================================================

def _pdf_footer(
    canvas,
    doc,
) -> None:
    canvas.saveState()

    width, _ = landscape(
        A4
    )

    canvas.setStrokeColor(
        colors.HexColor(
            "#E6DCEC"
        )
    )

    canvas.line(
        13 * mm,
        10 * mm,
        width -
        13 * mm,
        10 * mm,
    )

    canvas.setFont(
        "Helvetica",
        7,
    )

    canvas.setFillColor(
        colors.HexColor(
            f"#{MUTED}"
        )
    )

    canvas.drawString(
        13 * mm,
        6 * mm,
        (
            "Stock Intelligence "
            "- Executive Report"
        ),
    )

    canvas.drawRightString(
        width -
        13 * mm,
        6 * mm,
        f"Page {doc.page}",
    )

    canvas.restoreState()